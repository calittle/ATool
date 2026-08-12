#!/usr/bin/env python3
"""Render a generated Technical Design Markdown report to DOCX and PDF."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
WIDTH = Inches(6.5)


def set_font(run, name: str = "Calibri", size: float = 11, bold: bool | None = None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    tc_pr.append(fill)


def bookmark(paragraph, name: str) -> None:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", name)[:38]
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), safe)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_hyperlink(paragraph, text: str, target: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        hyperlink.set(qn("w:anchor"), re.sub(r"[^A-Za-z0-9_]", "_", target[1:])[:38])
    else:
        relationship_id = paragraph.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def append_inline(paragraph, value: str, *, size=11) -> None:
    value = re.sub(r'<a id="([^"]+)"></a>', '', value)
    token = re.compile(r"(\[[^\]]+\]\([^\)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    cursor = 0
    for match in token.finditer(value):
        if match.start() > cursor:
            set_font(paragraph.add_run(value[cursor:match.start()]), size=size)
        text = match.group(0)
        if text.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", text).groups()
            add_hyperlink(paragraph, label, url)
        elif text.startswith("`"):
            set_font(paragraph.add_run(text[1:-1]), name="Consolas", size=max(8, size - 1))
        elif text.startswith("**"):
            set_font(paragraph.add_run(text[2:-2]), size=size, bold=True)
        else:
            run = paragraph.add_run(text[1:-1])
            set_font(run, size=size)
            run.italic = True
        cursor = match.end()
    if cursor < len(value):
        set_font(paragraph.add_run(value[cursor:]), size=size)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    cols = len(rows[0])
    table = document.add_table(rows=0, cols=cols)
    table.autofit = False
    table.style = "Table Grid"
    widths = [int(9360 / cols)] * cols
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col, cell in enumerate(cells):
            cell.width = Inches(widths[col] / 1440)
            set_cell_margins(cell)
            if row_index == 0:
                shade(cell, "F2F4F7")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            anchor = re.search(r'<a id="([^"]+)"></a>', values[col])
            if anchor:
                bookmark(paragraph, anchor.group(1))
            append_inline(paragraph, values[col], size=8.5 if cols >= 4 else 9.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("Technical Design  |  "), size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def render(markdown_path: Path, docx_path: Path) -> None:
    document = Document()
    configure(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(.25)
                paragraph.paragraph_format.space_after = Pt(5)
                set_font(paragraph.add_run("\n".join(code_lines)), name="Consolas", size=8)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[-| ]+\|$", lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lines[index]):
                    rows.append([cell.strip() for cell in lines[index].strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level, text = len(heading.group(1)), heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run(text)
                set_font(run, size=22, bold=True, color=DARK_BLUE)
            else:
                paragraph = document.add_paragraph(style=f"Heading {level - 1}")
                append_inline(paragraph, text, size={2: 16, 3: 13, 4: 12}[level])
                bookmark(paragraph, re.sub(r"[^A-Za-z0-9_]", "_", text.lower().replace(" ", "-")))
            index += 1
            continue
        bullet = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet:
            depth, text = len(bullet.group(1)) // 2, bullet.group(2)
            paragraph = document.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            paragraph.paragraph_format.left_indent = Inches(.5 + .25 * depth)
            paragraph.paragraph_format.first_line_indent = Inches(-.25)
            paragraph.paragraph_format.space_after = Pt(2)
            append_inline(paragraph, text, size=9.5)
            index += 1
            continue
        if line.strip():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            append_inline(paragraph, line, size=10)
        index += 1
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Render TD Markdown as DOCX and PDF.")
    parser.add_argument("--markdown", required=True, type=Path)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--pdf", type=Path)
    args = parser.parse_args()
    render(args.markdown, args.docx)
    if args.pdf:
        soffice = shutil.which("soffice")
        if not soffice:
            raise RuntimeError("soffice is required to create a PDF")
        args.pdf.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(args.pdf.parent), str(args.docx)], check=True)
        produced = args.pdf.parent / f"{args.docx.stem}.pdf"
        if produced != args.pdf:
            produced.replace(args.pdf)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
